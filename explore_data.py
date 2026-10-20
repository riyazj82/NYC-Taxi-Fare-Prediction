import pandas as pd

# Load NYC Yellow Taxi dataset
file_path = "yellow_tripdata_2025-01.parquet"

df = pd.read_parquet(file_path)

# Display dataset information
print("Dataset loaded successfully!")

print("\nNumber of rows and columns:")
print(df.shape)

print("\nFirst 5 rows:")
print(df.head())

print("\nColumn names:")
print(df.columns.tolist())

print("\nMissing values:")
print(df.isnull().sum())

print("\nBasic statistics:")
print(df.describe())


from pathlib import Path
import numpy as np
import pandas as pd

from sklearn.model_selection import train_test_split
from sklearn.linear_model import LinearRegression
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import (
    mean_absolute_error,
    mean_squared_error,
    r2_score
)

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

# ==========================================
# 1. PROJECT CONFIGURATION
# ==========================================

BASE_DIR = Path(__file__).resolve().parent
DATA_FILE = BASE_DIR / "yellow_tripdata_2025-01.parquet"
OUTPUT_DIR = BASE_DIR / "results"
OUTPUT_DIR.mkdir(exist_ok=True)

RANDOM_STATE = 42
SAMPLE_SIZE = 50000

print("\nNYC TAXI FARE PREDICTION PROJECT")
print("=" * 55)

# ==========================================
# 2. LOAD DATA
# ==========================================

print("\nSTEP 1: Loading NYC Taxi dataset...")

if not DATA_FILE.exists():
    raise FileNotFoundError(
        f"Dataset not found: {DATA_FILE}"
    )

columns = [
    "trip_distance",
    "passenger_count",
    "PULocationID",
    "DOLocationID",
    "tpep_pickup_datetime",
    "fare_amount",
    "RatecodeID"
]

df = pd.read_parquet(DATA_FILE, columns=columns)

print("Dataset loaded successfully!")
print("Original dataset shape:", df.shape)

# ==========================================
# 3. DATA CLEANING
# ==========================================

print("\nSTEP 2: Cleaning dataset...")

df = df.dropna(subset=columns).copy()

df = df[
    (df["fare_amount"] > 0) &
    (df["fare_amount"] <= 200) &
    (df["trip_distance"] > 0) &
    (df["trip_distance"] <= 100) &
    (df["passenger_count"] >= 1) &
    (df["passenger_count"] <= 6)
].copy()

df["tpep_pickup_datetime"] = pd.to_datetime(
    df["tpep_pickup_datetime"]
)

df = df[
    (df["tpep_pickup_datetime"] >= "2025-01-01") &
    (df["tpep_pickup_datetime"] < "2025-02-01")
].copy()

print("Cleaned dataset shape:", df.shape)

if len(df) < 100:
    raise ValueError("Insufficient valid records for modeling.")

# ==========================================
# 4. FEATURE ENGINEERING
# ==========================================

print("\nSTEP 3: Creating features...")

df["pickup_hour"] = df["tpep_pickup_datetime"].dt.hour
df["pickup_day"] = df["tpep_pickup_datetime"].dt.dayofweek
df["is_weekend"] = (df["pickup_day"] >= 5).astype(int)

features = [
    "trip_distance",
    "passenger_count",
    "PULocationID",
    "DOLocationID",
    "pickup_hour",
    "pickup_day",
    "is_weekend",
    "RatecodeID"
]

# Use a reproducible sample for faster execution
if len(df) > SAMPLE_SIZE:
    df = df.sample(
        n=SAMPLE_SIZE,
        random_state=RANDOM_STATE
    )

X = df[features].copy()
y = df["fare_amount"].copy()

# Treat location and rate codes as categories
categorical_columns = [
    "PULocationID",
    "DOLocationID",
    "RatecodeID"
]

X = pd.get_dummies(
    X,
    columns=categorical_columns,
    dtype=int
)

print("Final sample size:", len(X))
print("Number of model features:", X.shape[1])

# ==========================================
# 5. TRAIN / TEST SPLIT
# ==========================================

print("\nSTEP 4: Splitting training and test data...")

X_train, X_test, y_train, y_test = train_test_split(
    X,
    y,
    test_size=0.20,
    random_state=RANDOM_STATE
)

print("Training records:", len(X_train))
print("Testing records:", len(X_test))

# ==========================================
# 6. TRAIN MODELS
# ==========================================

print("\nSTEP 5: Training Linear Regression...")

linear_model = LinearRegression()
linear_model.fit(X_train, y_train)

linear_predictions = linear_model.predict(X_test)

print("Linear Regression completed!")

print("\nSTEP 6: Training Random Forest...")

forest_model = RandomForestRegressor(
    n_estimators=50,
    max_depth=15,
    min_samples_leaf=3,
    random_state=RANDOM_STATE,
    n_jobs=-1
)

forest_model.fit(X_train, y_train)

forest_predictions = forest_model.predict(X_test)

print("Random Forest completed!")

# ==========================================
# 7. EVALUATE MODELS
# ==========================================

print("\nSTEP 7: Evaluating model performance...")

def evaluate_model(name, actual, predicted):
    return {
        "Model": name,
        "MAE": mean_absolute_error(actual, predicted),
        "RMSE": np.sqrt(
            mean_squared_error(actual, predicted)
        ),
        "R2": r2_score(actual, predicted)
    }

metrics = pd.DataFrame([
    evaluate_model(
        "Linear Regression",
        y_test,
        linear_predictions
    ),
    evaluate_model(
        "Random Forest",
        y_test,
        forest_predictions
    )
])

print("\nMODEL PERFORMANCE:")
print(metrics.to_string(index=False))

best_model = metrics.loc[
    metrics["MAE"].idxmin(), "Model"
]

print("\nBest model by MAE:", best_model)

# ==========================================
# 8. CREATE PREDICTIONS CSV
# ==========================================

print("\nSTEP 8: Exporting predictions...")

predictions = pd.DataFrame({
    "Actual_Fare": y_test.to_numpy(),
    "Linear_Prediction": linear_predictions,
    "Random_Forest_Prediction": forest_predictions
})

predictions["Linear_Absolute_Error"] = abs(
    predictions["Actual_Fare"] -
    predictions["Linear_Prediction"]
)

predictions["Forest_Absolute_Error"] = abs(
    predictions["Actual_Fare"] -
    predictions["Random_Forest_Prediction"]
)

predictions.to_csv(
    OUTPUT_DIR / "taxi_predictions.csv",
    index=False
)

metrics.to_csv(
    OUTPUT_DIR / "model_metrics.csv",
    index=False
)

# ==========================================
# 9. FEATURE IMPORTANCE
# ==========================================

print("\nSTEP 9: Calculating feature importance...")

importance = pd.DataFrame({
    "Feature": X.columns,
    "Importance": forest_model.feature_importances_
})

importance = importance.sort_values(
    "Importance",
    ascending=False
)

importance.to_csv(
    OUTPUT_DIR / "feature_importance.csv",
    index=False
)

# ==========================================
# 10. EXPORT EXCEL WORKBOOK
# ==========================================

print("\nSTEP 10: Creating Excel report...")

excel_file = OUTPUT_DIR / "NYC_Taxi_Model_Report.xlsx"

with pd.ExcelWriter(
    excel_file,
    engine="openpyxl"
) as writer:

    metrics.to_excel(
        writer,
        sheet_name="Model Comparison",
        index=False
    )

    predictions.head(1000).to_excel(
        writer,
        sheet_name="Predictions",
        index=False
    )

    importance.head(20).to_excel(
        writer,
        sheet_name="Feature Importance",
        index=False
    )

# ==========================================
# 11. FORMAT EXCEL AND ADD CHART
# ==========================================

wb = load_workbook(excel_file)

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            start_color="17365D",
            end_color="17365D",
            fill_type="solid"
        )

    for column in ws.columns:
        letter = column[0].column_letter
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        ws.column_dimensions[letter].width = min(
            max_length + 3, 32
        )

ws = wb["Model Comparison"]

chart = BarChart()
chart.title = "Model Mean Absolute Error"
chart.y_axis.title = "MAE (USD)"
chart.x_axis.title = "Model"

data = Reference(
    ws,
    min_col=2,
    min_row=1,
    max_row=3
)

categories = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=3
)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.width = 16
chart.height = 9

ws.add_chart(chart, "F2")

wb.save(excel_file)

# ==========================================
# 12. FINAL SUMMARY
# ==========================================

print("\n" + "=" * 55)
print("PROJECT COMPLETED SUCCESSFULLY")
print("=" * 55)

print("\nBest model:", best_model)

print("\nGenerated files:")

for file in sorted(OUTPUT_DIR.iterdir()):
    print(" -", file.name)

print("\nExcel report:", excel_file)

print("\nDone!")
